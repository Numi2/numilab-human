"""Append the measured tendon visual-QA page to the retained executive pack."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import zipfile

from pptx import Presentation
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[3]
HERE = Path(__file__).resolve().parent
REPORTS = Path('/Users/home/NumiSuiteReports')
PRIOR = REPORTS/'board-20260929-anatomy-verification'
OUT = REPORTS/'board-20260929-anatomy-verification-tendon'
PRIOR_STEM = 'Numi-Suite-Board-Visuals-anatomy-verification-20260929'
NEW_STEM = 'Numi-Suite-Board-Visuals-anatomy-verification-tendon-20260929'
SOFFICE = Path('/Users/home/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice')
PDFTOPPM = Path('/Users/home/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/pdftoppm')


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def render_page(pdf: Path, page: int, scale: int = 1600) -> bytes:
    return subprocess.run(
        [str(PDFTOPPM), '-f', str(page), '-l', str(page), '-scale-to',
         str(scale), '-singlefile', '-png', str(pdf)],
        check=True, capture_output=True,
    ).stdout


def main() -> None:
    assert PRIOR.is_dir() and not OUT.exists()
    prior_pptx = PRIOR/(PRIOR_STEM+'.pptx')
    prior_pdf = PRIOR/(PRIOR_STEM+'.pdf')
    assert prior_pptx.is_file() and prior_pdf.is_file()
    receipt = json.loads((HERE/'receipt-v1.json').read_text())
    assert receipt['changed_stable_ids'] == [7, 8]
    assert [row['new_exact_intersection_pairs'] for row in receipt['tendons']] == [0, 0]
    assert receipt['unchanged_surface_count'] == 148
    shutil.copytree(PRIOR, OUT)

    asset = OUT/'assets/executive-tendon-geometry.png'
    shutil.copy2(HERE/'executive-tendon-geometry.png', asset)
    evidence = OUT/'evidence/tendon-harmonic-boundary-20260929'
    evidence.mkdir()
    for name in ('receipt-v1.json', 'create_receipt.py', 'create_board_figure.py'):
        shutil.copy2(HERE/name, evidence/name)
    scan = ROOT/'Docs/media/muscle-surface-embeddedness-20260929/receipt-v3.json'
    shutil.copy2(scan, evidence/'muscle-surface-embeddedness-receipt-v3.json')
    shutil.copy2(ROOT/'Docs/TENDON_HARMONIC_BOUNDARY_20260929.md',
                 evidence/'TENDON_HARMONIC_BOUNDARY_20260929.md')

    p = Presentation(prior_pptx)
    assert len(p.slides) == 34
    w, h = p.slide_width, p.slide_height
    assert abs(w/h - 16/9) < .001
    slide = p.slides.add_slide(p.slide_layouts[6])
    slide.shapes.add_picture(str(asset), 0, 0, w, h)
    slide.notes_slide.notes_text_frame.text = (
        'Full 150-surface compiled FP32 exact-predicate comparison: right/left '
        'calcaneal tendon pairs 145/163 before, 0/0 now. Other 148 local '
        'geometries and body-binding bytes unchanged. Both tendon sheets '
        'remain open with 388/389 boundary edges and 1/2 locally reversed '
        'faces. Visual geometry only, no anatomical enthesis or loaded '
        'standing qualification. Figure SHA-256 ' + sha(asset)
    )
    new_pptx = OUT/(NEW_STEM+'.pptx')
    new_pdf = OUT/(NEW_STEM+'.pdf')
    p.save(new_pptx)
    with zipfile.ZipFile(prior_pptx) as before, zipfile.ZipFile(new_pptx) as after:
        parts = [name for name in before.namelist()
                 if name.startswith('ppt/slides/') or name.startswith('ppt/media/')]
        assert all(before.read(name) == after.read(name) for name in parts)
    subprocess.run(
        [str(SOFFICE),
         '-env:UserInstallation=file:///tmp/numi-board-tendon-geometry-20260929',
         '--headless', '--convert-to', 'pdf', '--outdir', str(OUT), str(new_pptx)],
        check=True, capture_output=True, timeout=240,
    )
    assert len(PdfReader(prior_pdf).pages) == 34
    assert len(PdfReader(new_pdf).pages) == 35
    for page in (1, 33, 34):
        assert render_page(prior_pdf, page, 800) == render_page(new_pdf, page, 800)
    (OUT/'review-page-35.png').write_bytes(render_page(new_pdf, 35))
    (OUT/(PRIOR_STEM+'.pptx')).unlink()
    (OUT/(PRIOR_STEM+'.pdf')).unlink()
    (OUT/'README.md').write_text(
        '# Numi Suite executive board visuals - anatomy verification update\n\n'
        '35 slides. Slide 34 shows current projected patellar anteriority; '
        'slide 35 shows the compiled calcaneal tendon geometry delta. '
        'The preceding 34 PPTX slide/media parts are byte-identical, and '
        'sampled PDF pages 1, 33 and 34 are pixel-identical to the prior pack.\n\n'
        'The included 10-second standing clip is the September 28 native run. '
        'It predates the patellar audit and tendon repair, so it is historical '
        'simulation footage, not a standing run of the current anatomy. '
        'The visual repairs do not qualify clinical anatomy, tendon force, '
        'loaded contact or sustained standing.\n\n'
        'BodyParts3D CC-BY-SA 2.1 Japan; Z-Anatomy CC-BY-SA 4.0. '
        'Receipts and source attribution are in evidence/.\n'
    )
    historical = OUT/'Numi-Human-10-second-standing-simulation.mp4'
    (OUT/'manifest.json').write_text(json.dumps({
        'schema': 'numi.suite.board-anatomy-verification-tendon.v1',
        'date': '2026-09-29', 'slide_count': 35,
        'previous_34_slide_parts_and_media_byte_identical': True,
        'prior_pdf_pixel_identical_pages': [1, 33, 34],
        'prior_manifest_sha256': sha(PRIOR/'manifest.json'),
        'patella_figure_sha256': sha(OUT/'assets/executive-patellar-orientation.png'),
        'tendon_figure_sha256': sha(asset),
        'tendon_delta_receipt_sha256': sha(evidence/'receipt-v1.json'),
        'tendon_full_scan_sha256': sha(evidence/'muscle-surface-embeddedness-receipt-v3.json'),
        'pptx_sha256': sha(new_pptx), 'pdf_sha256': sha(new_pdf),
        'historical_video': {'date': '2026-09-28', 'sha256': sha(historical),
                             'predates_current_anatomy': True},
    }, indent=2, sort_keys=True)+'\n')
    print(json.dumps({'pptx': str(new_pptx), 'pdf': str(new_pdf),
                      'slides': 35, 'review': str(OUT/'review-page-35.png')}))


if __name__ == '__main__':
    main()
