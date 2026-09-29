import json,hashlib,shutil,zipfile
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from PIL import Image
base=Path('/Users/home/NumiSuiteReports');old=base/'board-20260929-organ-families';out=base/'board-20260929-whole-visceral';out.mkdir()
for name in ['assets','evidence']:shutil.copytree(old/name,out/name)
video='Numi-Human-10-second-standing-simulation.mp4';shutil.copy2(old/video,out/video)
source=old/'Numi-Suite-Board-Visuals-organ-family-update.pptx';r=Presentation(source);W,H=r.slide_width,r.slide_height
media=Path('/Users/home/numilab-human/Docs/media/whole-visceral-coverage-20260929')
slides=[('executive-brain-eye-coverage.png','Actual native source-family views. The brain family has 59 members: 54 neural regions in the displayed mask and five separately inspected ventricular-system regions. The complete eye families have 45 unique members, preserving 14 ocular muscle references, 23 ocular regions, six lacrimal-duct segments and two immaterial lacrimal lakes. Opaque source surfaces may occlude interior members; packaged count does not prove each surface has visible pixels. Seven cranial bone meshes are deliberately omitted from these head inspection packets while retaining mandible and other bones; the bone payload is unchanged. All 579 anatomy surfaces remain in every packet. Eight skull registrations agree algebraically with the common atlas frame; this is not intracranial containment or clinical placement proof. Source head and neck are fixed descendants of torso. Independent cervical and eye motion remain open.'),('executive-visceral-coverage.png','Adding 192 pinned source members extends 387 to 579 surfaces while preserving every previous vertex, record and triangle. All 571 unique members of 46 declared source families are represented. Small and large intestines have 63 unique members including the shared ileocecal junction; rectum is also shared and explicitly bound to pelvis. Urinary and male reproductive structures plus adrenal, salivary and lacrimal glands are included. Source families can contain overlapping aggregate/descendant representations and cannot be summed as disjoint tissue volume. The native gut front image is cropped for the slide without changing geometry. Fifteen added meshes retain source topology defects. Source/body-frame geometry passes in raw rest, projected neutral and coupled torso pose, with maximum added native/source position error 0.1832 micrometres under the unchanged 20-micrometre gate. Connected lumens, physical tissue/fluid volume, subject/clinical anatomy, tissue deformation and mechanics remain open. The earlier standing video remains dated 28 September and predates this anatomy payload.')]
for name,notes in slides:
 p=media/name;shutil.copy2(p,out/'assets'/name);s=r.slides.add_slide(r.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('101C2B')
 iw,ih=Image.open(p).size;scale=min(W/iw,H/ih);width,height=round(iw*scale),round(ih*scale)
 s.shapes.add_picture(str(p),(W-width)//2,(H-height)//2,width,height);s.notes_slide.notes_text_frame.text=notes+'\nSource image SHA256 '+hashlib.sha256(p.read_bytes()).hexdigest()
pptx=out/'Numi-Suite-Board-Visuals-whole-visceral-update.pptx';r.save(pptx)
with zipfile.ZipFile(source) as a,zipfile.ZipFile(pptx) as b:
 parts=[n for n in a.namelist() if n.startswith('ppt/slides/') or n.startswith('ppt/media/')]
 assert all(a.read(n)==b.read(n) for n in parts)
(out/'source-slide-parity.json').write_text(json.dumps({'source_deck':str(source),'source_deck_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'original_slide_count':28,'new_slide_count':len(r.slides),'original_slide_parts_and_media_byte_identical':True,'checked_zip_parts':len(parts)},indent=2)+'\n')
(out/'README.md').write_text('''# Numi Suite executive-board visual pack

30 slides: the preceding 28-slide suite pack, plus two current brain/eye and
visceral source-coverage slides. All preceding slide parts and media remain
byte-identical; their original dates and evidence limits remain.

The Human composite now contains 579 surfaces with all 571 unique members of
46 declared source families represented. Adding 192 source members preserves
the original 387 surfaces byte for byte. Independent source/body-frame audits
cover every surface in rest, neutral and torso-flexed poses. This is source
membership and native reference geometry, not whole-Human or clinical anatomy.

Fifteen new meshes retain source topology defects. Source head/neck lack
independent cervical/eye motion. Tissue volume, connected lumens, clinical
placement, subject anatomy, deformation and organ mechanics remain open.

The included 10-second standing video is the 28 September native run. It
predates these anatomy changes and is not a standing run of this payload.

BodyParts3D - The Database Center for Life Science - CC-BY-SA 2.1 Japan;
Z-Anatomy - The libre 3D atlas of anatomy - CC-BY-SA 4.0. Full pinned-source
and derivative attribution is retained with the evidence.
''')
print(pptx)
