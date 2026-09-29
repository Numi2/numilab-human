import hashlib,json,shutil,zipfile
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from PIL import Image
base=Path('/Users/home/NumiSuiteReports');old=base/'board-20260929-lung-envelopes';out=base/'board-20260929-organ-families';out.mkdir()
for name in ['assets','evidence']:
 if (old/name).exists():shutil.copytree(old/name,out/name)
video='Numi-Human-10-second-standing-simulation.mp4';shutil.copy2(old/video,out/video)
source=old/'Numi-Suite-Board-Visuals-lung-envelope-update.pptx';r=Presentation(source);W,H=r.slide_width,r.slide_height
media=Path('/Users/home/numilab-human/Docs/media/organ-family-coverage-20260929')
slides=[('executive-organ-family-progress.png','The existing 18-region atlas graph contains 378 unique members. All are now represented by adding 77 omitted source members: 18 material organ components, 40 vessel segments, 15 ducts and four cardiac cavity references. The previous 310 surfaces are byte-identical within the new 387-surface composite. Family counts overlap. All 387 surfaces pass source-to-native geometry checks in three poses and four viewing profiles; maximum added position error is 0.120 micrometres. Complete declared source membership does not establish exhaustive whole-body or clinical anatomy.'),('executive-organ-family-types.png','Actual native projected-neutral camera captures of the same complete payload, with separate cavity-reference and duct masks. Images are cropped for the slide without changing geometry. Three added members retain duplicate-face or vertex-link defects: FJ2404, FJ2405 and FJ2434. The existing source right-atrium/right-ventricle overlap remains; these are reference surfaces, not admitted physical blood volumes. Aggregate/descendant atlas representations can overlap. Connected lumens, tissue volumes, clinical registration and organ mechanics remain unqualified. Retain BodyParts3D and Z-Anatomy source attribution.')]
for name,notes in slides:
 p=media/name;shutil.copy2(p,out/'assets'/name)
 s=r.slides.add_slide(r.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('101C2B')
 iw,ih=Image.open(p).size;scale=min(W/iw,H/ih);width,height=round(iw*scale),round(ih*scale)
 s.shapes.add_picture(str(p),(W-width)//2,(H-height)//2,width,height)
 s.notes_slide.notes_text_frame.text=notes+'\nSource image SHA256 '+hashlib.sha256(p.read_bytes()).hexdigest()
pptx=out/'Numi-Suite-Board-Visuals-organ-family-update.pptx';r.save(pptx)
with zipfile.ZipFile(source) as a,zipfile.ZipFile(pptx) as b:
 unchanged=[n for n in a.namelist() if n.startswith('ppt/slides/') or n.startswith('ppt/media/')]
 assert all(a.read(n)==b.read(n) for n in unchanged)
(out/'source-slide-parity.json').write_text(json.dumps({'source_deck':str(source),'source_deck_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'original_slide_count':26,'new_slide_count':len(r.slides),'original_slide_parts_and_media_byte_identical':True,'checked_zip_parts':len(unchanged)},indent=2)+'\n')
(out/'README.md').write_text("""# Numi Suite executive-board visual pack

28 slides: the original suite evidence and earlier 29 September knee, skin,
branch and lung-envelope additions, followed by two current source-family
slides. Slides 1-26 retain their original dates and evidence limits.

All 378 unique members of the existing 18-region atlas graph are represented.
Adding 77 omitted members produces a 387-surface composite while retaining
the original 310 surfaces byte for byte. All 387 surfaces pass source geometry
audits in three poses and four masks. Shared family members and overlapping
aggregate/descendant atlas representations must not be summed as disjoint
tissue. Three added members retain source defects; the known right-heart
cavity overlap remains. Clinical anatomy, exhaustive whole-body coverage,
connected lumens, tissue volumes and mechanics remain open.

The included 10-second standing video remains the 28 September native run.
It predates the later anatomy repairs and is not a current standing run of
the 387-surface payload.

Source-derived images retain the existing credits: BodyParts3D - The Database
Center for Life Science - CC-BY-SA 2.1 Japan; Z-Anatomy - The libre 3D atlas of
anatomy - CC-BY-SA 4.0. The complete source/derivative attribution is included
with the evidence.
""")
print(pptx)
