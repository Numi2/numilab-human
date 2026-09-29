"""Append the source-bound muscle tip visual repair to the executive deck."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import zipfile

from pptx import Presentation
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[3]
HERE = Path(__file__).resolve().parent
REPORTS = Path('/Users/home/NumiSuiteReports')
PRIOR = REPORTS/'board-20260929-anatomy-verification-expanded-organs'
OUT = REPORTS/'board-20260929-anatomy-verification-muscle-tip'
PRIOR_STEM = 'Numi-Suite-Board-Visuals-anatomy-verification-expanded-organs-20260929'
NEW_STEM = 'Numi-Suite-Board-Visuals-anatomy-verification-muscle-tip-20260929'
OVERRIDE = Path('/Users/home/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override')


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def render_page(pdf: Path, page: int) -> bytes:
    return subprocess.run(
        [str(OVERRIDE/'pdftoppm'), '-f', str(page), '-l', str(page),
         '-scale-to', '800', '-singlefile', '-png', str(pdf)],
        check=True, capture_output=True,
    ).stdout


def main() -> None:
    assert PRIOR.is_dir() and not OUT.exists()
    prior_pptx = PRIOR/(PRIOR_STEM+'.pptx')
    prior_pdf = PRIOR/(PRIOR_STEM+'.pdf')
    receipt = json.loads((HERE/'receipt-v1.json').read_text())
    assert receipt['changed_vertex_record_count'] == 32
    assert receipt['single_embedded_muscle_surfaces_after'] == 54
    assert len(receipt['changed_surfaces']) == 2
    shutil.copytree(PRIOR, OUT)
    asset = OUT/'assets/executive-biceps-tip-untangle.png'
    shutil.copy2(HERE/'executive-biceps-tip-untangle.png', asset)
    evidence = OUT/'evidence/biceps-tip-visual-untangle-20260929'
    evidence.mkdir()
    for name in ('receipt-v1.json', 'create_board_figure.py'):
        shutil.copy2(HERE/name, evidence/name)
    for source in ('muscle_tip_visual_untangle.py', 'muscle_tip_visual_delta.py'):
        shutil.copy2(ROOT/'src/numilab_human'/source, evidence/source)
    shutil.copy2(ROOT/'Docs/BICEPS_SHORT_HEAD_VISUAL_UNTANGLE_20260929.md',
                 evidence/'BICEPS_SHORT_HEAD_VISUAL_UNTANGLE_20260929.md')
    shutil.copy2(ROOT/'Docs/media/muscle-surface-embeddedness-20260929/receipt-v4.json',
                 evidence/'muscle-surface-embeddedness-v4.json')
    shutil.copy2(ROOT/'Build/biceps-tip-visual-untangle-20260929/matched/bodyparts3d-myosim-fullbody-muscle-surfaces.manifest.json',
                 evidence/'compiled-muscle-surface-manifest.json')

    deck = Presentation(prior_pptx)
    assert len(deck.slides) == 37
    width, height = deck.slide_width, deck.slide_height
    assert abs(width/height-16/9) < .001
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(asset), 0, 0, width, height)
    slide.notes_slide.notes_text_frame.text = (
        'The pinned right and left biceps femoris short-head source members '
        'each have three exact tip intersections. Two explicitly derived '
        'visual candidates now have zero in the compiled Float32 packet. '
        'Single embedded muscle visual candidates rise from 52 to 54 of '
        '148. Across the full 150-surface packet, headers, records, '
        'bindings, indices, skinning weights, and all unrelated vertex '
        'records are byte-identical. Only twelve positions and adjacent '
        'normals change, with about 0.065 mm maximum displacement. This '
        'is not clinical anatomy or loaded muscle mechanics. Figure SHA-256 '
        + sha(asset)
    )
    pptx = OUT/(NEW_STEM+'.pptx')
    pdf = OUT/(NEW_STEM+'.pdf')
    deck.save(pptx)
    with zipfile.ZipFile(prior_pptx) as before, zipfile.ZipFile(pptx) as after:
        parts = [name for name in before.namelist()
                 if name.startswith('ppt/slides/') or name.startswith('ppt/media/')]
        assert all(before.read(name) == after.read(name) for name in parts)
    subprocess.run(
        [str(OVERRIDE/'soffice'),
         '-env:UserInstallation=file:///tmp/numi-board-muscle-tip-20260929',
         '--headless', '--convert-to', 'pdf', '--outdir', str(OUT), str(pptx)],
        check=True, capture_output=True, timeout=240,
    )
    assert len(PdfReader(prior_pdf).pages) == 37
    assert len(PdfReader(pdf).pages) == 38
    for page in (1, 36, 37):
        assert render_page(prior_pdf, page) == render_page(pdf, page)
    (OUT/'review-page-38.png').write_bytes(render_page(pdf, 38))
    (OUT/(PRIOR_STEM+'.pptx')).unlink()
    (OUT/(PRIOR_STEM+'.pdf')).unlink()
    (OUT/'README.md').write_text(
        '# Numi Suite executive board visuals - bilateral muscle tip update\n\n'
        '38 slides. Slide 38 records the source-bound visual-only biceps '
        'femoris tip repair. The preceding 37 PPTX slide/media parts are '
        'byte-identical, and sampled PDF pages 1, 36 and 37 render '
        'pixel-identically to the prior deck. Clinical muscle anatomy, '
        'contact, material and force transfer are still open.\n\n'
        'The included 10-second standing clip is the September 28 native '
        'simulation run and predates these anatomy changes.\n\n'
        'BodyParts3D CC-BY-SA 2.1 Japan; Z-Anatomy CC-BY-SA 4.0. '
        'Receipts and attribution are in evidence/.\n'
    )
    historical = OUT/'Numi-Human-10-second-standing-simulation.mp4'
    (OUT/'manifest.json').write_text(json.dumps({
        'schema': 'numi.suite.board-anatomy-verification-muscle-tip.v1',
        'date': '2026-09-29', 'slide_count': 38,
        'previous_37_slide_parts_and_media_byte_identical': True,
        'prior_pdf_pixel_identical_pages': [1, 36, 37],
        'prior_manifest_sha256': sha(PRIOR/'manifest.json'),
        'muscle_tip_figure_sha256': sha(asset),
        'muscle_tip_receipt_sha256': sha(evidence/'receipt-v1.json'),
        'muscle_full_census_sha256': sha(evidence/'muscle-surface-embeddedness-v4.json'),
        'pptx_sha256': sha(pptx), 'pdf_sha256': sha(pdf),
        'historical_video': {'date': '2026-09-28', 'sha256': sha(historical),
                             'predates_current_anatomy': True},
    }, indent=2, sort_keys=True)+'\n')
    print(json.dumps({'slides': 38, 'pptx': str(pptx),
                      'pdf': str(pdf), 'review': str(OUT/'review-page-38.png')}))


if __name__ == '__main__':
    main()
