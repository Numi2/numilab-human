"""Append the failed internal-organ separation gate to the board visual pack."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import zipfile

from pptx import Presentation
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[3]
HERE = Path(__file__).resolve().parent
REPORTS = Path('/Users/home/NumiSuiteReports')
PRIOR = REPORTS/'board-20260929-anatomy-verification-tendon'
OUT = REPORTS/'board-20260929-anatomy-verification-organs'
PRIOR_STEM = 'Numi-Suite-Board-Visuals-anatomy-verification-tendon-20260929'
NEW_STEM = 'Numi-Suite-Board-Visuals-anatomy-verification-organs-20260929'
SOFFICE = Path('/Users/home/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice')
PDFTOPPM = Path('/Users/home/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/pdftoppm')


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def render_page(pdf: Path, page: int, scale: int = 1600) -> bytes:
    return subprocess.run(
        [str(PDFTOPPM), '-f', str(page), '-l', str(page), '-scale-to',
         str(scale), '-singlefile', '-png', str(pdf)],
        check=True, capture_output=True,
    ).stdout


def main() -> None:
    assert PRIOR.is_dir() and not OUT.exists()
    prior_pptx = PRIOR/(PRIOR_STEM+'.pptx')
    prior_pdf = PRIOR/(PRIOR_STEM+'.pdf')
    receipt = json.loads((HERE/'receipt-v1.json').read_text())
    witness = json.loads((HERE/'transverse-witness-v1.json').read_text())
    assert receipt['status'] == 'failed_cross_surface_separation'
    assert receipt['pair_count'] == 10 and receipt['clear_pair_count'] == 6
    assert len(receipt['failed_pair_names']) == len(witness['witnesses']) == 4
    shutil.copytree(PRIOR, OUT)
    asset = OUT/'assets/executive-organ-separation.png'
    shutil.copy2(HERE/'executive-organ-separation.png', asset)
    evidence = OUT/'evidence/abdominal-organ-separation-20260929'
    evidence.mkdir()
    for name in ('receipt-v1.json', 'transverse-witness-v1.json',
                 'independent_transverse_witness.py', 'create_board_figure.py'):
        shutil.copy2(HERE/name, evidence/name)
    shutil.copy2(ROOT/'Docs/ABDOMINAL_ORGAN_SEPARATION_20260929.md',
                 evidence/'ABDOMINAL_ORGAN_SEPARATION_20260929.md')

    p = Presentation(prior_pptx)
    assert len(p.slides) == 35
    width, height = p.slide_width, p.slide_height
    assert abs(width/height - 16/9) < .001
    slide = p.slides.add_slide(p.slide_layouts[6])
    slide.shapes.add_picture(str(asset), 0, 0, width, height)
    slide.notes_slide.notes_text_frame.text = (
        'Five individually closed and exactly embedded source-named organ '
        'representations were tested pairwise in the selected 602-surface '
        'native anatomy packet. Six of ten pairs are separate. Four have '
        '182 exact crossing triangle pairs total, with independent strict '
        'transverse witnesses. This fails disjoint organ-domain admission; '
        'there was no source movement or clinical registration. Figure SHA-256 '
        + sha(asset)
    )
    new_pptx = OUT/(NEW_STEM+'.pptx')
    new_pdf = OUT/(NEW_STEM+'.pdf')
    p.save(new_pptx)
    with zipfile.ZipFile(prior_pptx) as before, zipfile.ZipFile(new_pptx) as after:
        parts = [name for name in before.namelist()
                 if name.startswith('ppt/slides/') or name.startswith('ppt/media/')]
        assert all(before.read(name) == after.read(name) for name in parts)
    subprocess.run(
        [str(SOFFICE),
         '-env:UserInstallation=file:///tmp/numi-board-organ-separation-20260929',
         '--headless', '--convert-to', 'pdf', '--outdir', str(OUT), str(new_pptx)],
        check=True, capture_output=True, timeout=240,
    )
    assert len(PdfReader(prior_pdf).pages) == 35
    assert len(PdfReader(new_pdf).pages) == 36
    for page in (1, 34, 35):
        assert render_page(prior_pdf, page, 800) == render_page(new_pdf, page, 800)
    (OUT/'review-page-36.png').write_bytes(render_page(new_pdf, 36))
    (OUT/(PRIOR_STEM+'.pptx')).unlink()
    (OUT/(PRIOR_STEM+'.pdf')).unlink()
    (OUT/'README.md').write_text(
        '# Numi Suite executive board visuals - organ separation update\n\n'
        '36 slides. Slide 36 records the failed exact five-organ separation '
        'gate. The preceding 35 PPTX slide/media parts are byte-identical, '
        'and sampled PDF pages 1, 34 and 35 render pixel-identically to the '
        'prior pack. The measured result is six separate organ pairs and four '
        'crossing pairs; no clinical anatomy or organ mechanics is admitted.\n\n'
        'The included 10-second standing clip is the September 28 native run '
        'and predates the recent anatomy changes. It is historical simulation '
        'footage, not a current-anatomy standing result.\n\n'
        'BodyParts3D CC-BY-SA 2.1 Japan; Z-Anatomy CC-BY-SA 4.0. '
        'Receipts and attribution are in evidence/.\n'
    )
    historical = OUT/'Numi-Human-10-second-standing-simulation.mp4'
    (OUT/'manifest.json').write_text(json.dumps({
        'schema': 'numi.suite.board-anatomy-verification-organs.v1',
        'date': '2026-09-29', 'slide_count': 36,
        'previous_35_slide_parts_and_media_byte_identical': True,
        'prior_pdf_pixel_identical_pages': [1, 34, 35],
        'prior_manifest_sha256': sha(PRIOR/'manifest.json'),
        'organ_separation_status': receipt['status'],
        'organ_figure_sha256': sha(asset),
        'organ_receipt_sha256': sha(evidence/'receipt-v1.json'),
        'independent_witness_sha256': sha(evidence/'transverse-witness-v1.json'),
        'pptx_sha256': sha(new_pptx), 'pdf_sha256': sha(new_pdf),
        'historical_video': {'date': '2026-09-28', 'sha256': sha(historical),
                             'predates_current_anatomy': True},
    }, indent=2, sort_keys=True)+'\n')
    print(json.dumps({'slides': 36, 'pptx': str(new_pptx),
                      'pdf': str(new_pdf), 'review': str(OUT/'review-page-36.png')}))


if __name__ == '__main__':
    main()
