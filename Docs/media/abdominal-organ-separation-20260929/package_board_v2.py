"""Append the expanded eight-organ gate to the existing executive board pack."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import zipfile

from pptx import Presentation
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[3]
HERE = Path(__file__).resolve().parent
REPORTS = Path('/Users/home/NumiSuiteReports')
PRIOR = REPORTS/'board-20260929-anatomy-verification-organs'
OUT = REPORTS/'board-20260929-anatomy-verification-expanded-organs'
PRIOR_STEM = 'Numi-Suite-Board-Visuals-anatomy-verification-organs-20260929'
NEW_STEM = 'Numi-Suite-Board-Visuals-anatomy-verification-expanded-organs-20260929'
OVERRIDE = Path('/Users/home/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override')


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def render_page(pdf: Path, page: int) -> bytes:
    return subprocess.run(
        [str(OVERRIDE/'pdftoppm'), '-f', str(page), '-l', str(page),
         '-scale-to', '800', '-singlefile', '-png', str(pdf)],
        check=True, capture_output=True,
    ).stdout


def main() -> None:
    assert PRIOR.is_dir() and not OUT.exists()
    prior_pptx = PRIOR/(PRIOR_STEM+'.pptx')
    prior_pdf = PRIOR/(PRIOR_STEM+'.pdf')
    receipt = json.loads((HERE/'receipt-v2.json').read_text())
    witness = json.loads((HERE/'transverse-witness-v2.json').read_text())
    assert receipt['status'] == 'failed_cross_surface_separation'
    assert receipt['pair_count'] == 28 and receipt['clear_pair_count'] == 21
    assert len(receipt['failed_pair_names']) == len(witness['witnesses']) == 7
    assert witness['primary_receipt_sha256'] == sha(HERE/'receipt-v2.json')
    shutil.copytree(PRIOR, OUT)
    asset = OUT/'assets/executive-organ-separation-v2.png'
    shutil.copy2(HERE/'executive-organ-separation-v2.png', asset)
    evidence = OUT/'evidence/expanded-abdominal-organ-separation-20260929'
    evidence.mkdir()
    for name in ('receipt-v2.json', 'transverse-witness-v2.json',
                 'independent_transverse_witness_v2.py', 'create_board_figure_v2.py'):
        shutil.copy2(HERE/name, evidence/name)
    shutil.copy2(ROOT/'src/numilab_human/expanded_abdominal_organ_separation.py',
                 evidence/'expanded_abdominal_organ_separation.py')
    shutil.copy2(ROOT/'Docs/EXPANDED_ABDOMINAL_ORGAN_SEPARATION_20260929.md',
                 evidence/'EXPANDED_ABDOMINAL_ORGAN_SEPARATION_20260929.md')

    p = Presentation(prior_pptx)
    assert len(p.slides) == 36
    width, height = p.slide_width, p.slide_height
    assert abs(width/height - 16/9) < .001
    slide = p.slides.add_slide(p.slide_layouts[6])
    slide.shapes.add_picture(str(asset), 0, 0, width, height)
    slide.notes_slide.notes_text_frame.text = (
        'Eight independently source-named, individually closed and exactly '
        'embedded organ surfaces were checked in one native Abdomen frame. '
        'Twenty-one of 28 pairs are separate; seven contain 814 exact '
        'intersecting triangle pairs, each with an independent strict '
        'transverse witness. This fails independent organ-domain admission. '
        'No atlas source geometry was moved and no clinical registration is '
        'claimed. Figure SHA-256 ' + sha(asset)
    )
    new_pptx = OUT/(NEW_STEM+'.pptx')
    new_pdf = OUT/(NEW_STEM+'.pdf')
    p.save(new_pptx)
    with zipfile.ZipFile(prior_pptx) as before, zipfile.ZipFile(new_pptx) as after:
        parts = [name for name in before.namelist()
                 if name.startswith('ppt/slides/') or name.startswith('ppt/media/')]
        assert all(before.read(name) == after.read(name) for name in parts)
    subprocess.run(
        [str(OVERRIDE/'soffice'),
         '-env:UserInstallation=file:///tmp/numi-board-expanded-organs-20260929',
         '--headless', '--convert-to', 'pdf', '--outdir', str(OUT), str(new_pptx)],
        check=True, capture_output=True, timeout=240,
    )
    assert len(PdfReader(prior_pdf).pages) == 36
    assert len(PdfReader(new_pdf).pages) == 37
    for page in (1, 35, 36):
        assert render_page(prior_pdf, page) == render_page(new_pdf, page)
    (OUT/'review-page-37.png').write_bytes(render_page(new_pdf, 37))
    (OUT/(PRIOR_STEM+'.pptx')).unlink()
    (OUT/(PRIOR_STEM+'.pdf')).unlink()
    (OUT/'README.md').write_text(
        '# Numi Suite executive board visuals - expanded organ audit\n\n'
        '37 slides. Slide 37 records the failed eight-organ exact separation '
        'gate. The preceding 36 PPTX slide and media parts are byte-identical, '
        'and sampled PDF pages 1, 35 and 36 render pixel-identically to the '
        'prior pack. Twenty-one pairs separate and seven pairs cross. '
        'No clinical anatomy or organ mechanics is admitted.\n\n'
        'The included 10-second standing clip is the September 28 native run '
        'and predates the recent anatomy changes. It is historical simulation '
        'footage, not a current-anatomy standing result.\n\n'
        'BodyParts3D CC-BY-SA 2.1 Japan; Z-Anatomy CC-BY-SA 4.0. '
        'Receipts and attribution are in evidence/.\n'
    )
    historical = OUT/'Numi-Human-10-second-standing-simulation.mp4'
    (OUT/'manifest.json').write_text(json.dumps({
        'schema': 'numi.suite.board-anatomy-verification-expanded-organs.v1',
        'date': '2026-09-29', 'slide_count': 37,
        'previous_36_slide_parts_and_media_byte_identical': True,
        'prior_pdf_pixel_identical_pages': [1, 35, 36],
        'prior_manifest_sha256': sha(PRIOR/'manifest.json'),
        'organ_separation_status': receipt['status'],
        'organ_figure_sha256': sha(asset),
        'organ_receipt_sha256': sha(evidence/'receipt-v2.json'),
        'independent_witness_sha256': sha(evidence/'transverse-witness-v2.json'),
        'pptx_sha256': sha(new_pptx), 'pdf_sha256': sha(new_pdf),
        'historical_video': {'date': '2026-09-28', 'sha256': sha(historical),
                             'predates_current_anatomy': True},
    }, indent=2, sort_keys=True)+'\n')
    print(json.dumps({'slides': 37, 'pptx': str(new_pptx),
                      'pdf': str(new_pdf), 'review': str(OUT/'review-page-37.png')}))


if __name__ == '__main__':
    main()
